import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def generar_xlsx(titulo: str, columnas: list, filas: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (titulo or 'Reporte')[:31]

    ws.append(columnas)
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='0C0E12')
        celda.alignment = Alignment(vertical='center')

    for fila in filas:
        ws.append(fila)

    for i, columna in enumerate(columnas, start=1):
        largo = len(str(columna))
        for fila in filas:
            largo = max(largo, len(str(fila[i - 1])))
        ws.column_dimensions[get_column_letter(i)].width = min(largo + 4, 42)

    ws.freeze_panes = 'A2'

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
